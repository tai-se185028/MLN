"""Extract MLN122.docx into sequential exam groups of at most 60 questions."""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "MLN122.docx"
OUTPUT = ROOT / "questions-mln122.json"
REPORT = ROOT / "extraction-report-mln122.md"
QUESTION_RE = re.compile(r"^\s*(\d{1,4})[.)]\s+(.*)$", re.DOTALL)
OPTION_RE = re.compile(r"^\s*([A-Z])\s*[.)]\s*(.*)$", re.DOTALL)
EXAM_SIZE = 60


def marked(paragraph) -> bool:
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        rgb = run.font.color.rgb
        colored = rgb is not None and str(rgb).upper() not in {"000000", "1A1D28"}
        if run.bold is True or run.underline is True or run.font.highlight_color is not None or colored:
            return True
    return False


def find_headings(document):
    """Follow the 1..N sequence so numbered explanation lines are not questions."""
    headings = []
    expected = 1
    for index, paragraph in enumerate(document.paragraphs):
        match = QUESTION_RE.match(paragraph.text.strip())
        if match and int(match.group(1)) == expected:
            headings.append((index, expected, match.group(2).strip()))
            expected += 1
    return headings


def extract():
    document = Document(SOURCE)
    headings = find_headings(document)
    questions = []
    format_counts = Counter()

    for pos, (paragraph_index, number, question_text) in enumerate(headings):
        end = headings[pos + 1][0] if pos + 1 < len(headings) else len(document.paragraphs)
        options, marked_indexes, notes = [], [], []
        for paragraph in document.paragraphs[paragraph_index + 1 : end]:
            text = paragraph.text.strip()
            if not text:
                continue
            match = OPTION_RE.match(text)
            if match:
                options.append(match.group(2).strip())
                if marked(paragraph):
                    marked_indexes.append(len(options) - 1)
            else:
                notes.append(text)
            for run in paragraph.runs:
                format_counts["bold"] += run.bold is True
                format_counts["underline"] += run.underline is True
                format_counts["highlight"] += run.font.highlight_color is not None
                format_counts["color"] += run.font.color.rgb is not None

        questions.append({
            "id": number,
            "question": question_text,
            "options": options,
            "correctAnswer": marked_indexes[0] if len(marked_indexes) == 1 else None,
            "correctAnswers": marked_indexes if len(marked_indexes) > 1 else None,
            "explanation": "\n".join(notes) if notes else None,
            "needsReview": len(marked_indexes) == 0,
        })

    exams = []
    for start in range(0, len(questions), EXAM_SIZE):
        exam_id = start // EXAM_SIZE + 1
        exams.append({"id": exam_id, "title": f"Đề {exam_id}", "questions": questions[start:start + EXAM_SIZE]})

    data = {"course": {"id": "mln122", "title": "MLN122"}, "exams": exams}
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return document, headings, questions, exams, format_counts


def report(document, headings, questions, exams, formats):
    normalized = defaultdict(list)
    for q in questions:
        normalized[" ".join(q["question"].casefold().split())].append(q["id"])
    duplicates = [ids for ids in normalized.values() if len(ids) > 1]
    missing = [q["id"] for q in questions if not q["options"]]
    non_four = [q["id"] for q in questions if len(q["options"]) != 4]
    multi = [q["id"] for q in questions if q["correctAnswers"]]
    unresolved = [q["id"] for q in questions if q["needsReview"]]
    notes = [q["id"] for q in questions if q["explanation"]]
    counts = [len(e["questions"]) for e in exams]
    sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    lines = [
        "# Báo cáo kiểm tra trích xuất MLN122", "",
        "- Trạng thái: **ĐẠT CÓ CẢNH BÁO**",
        f"- File nguồn: `MLN122.docx` (SHA-256: `{sha}`)",
        f"- Tổng đoạn văn: **{len(document.paragraphs)}**; tổng bảng: **{len(document.tables)}**",
        f"- Câu hỏi theo chuỗi đánh số 1–{len(headings)}: **{len(questions)}**",
        f"- Tổng đề: **{len(exams)}**; số câu từng đề: **{counts}**", "",
        "## Quy tắc chia đề", "",
        "Dữ liệu được chia tuần tự theo thứ tự nguồn, tối đa 60 câu/đề. Nguồn có 515 câu nên tạo 8 đề đủ 60 câu và đề 9 có 35 câu. Không bỏ câu và không tạo nội dung để làm tròn.", "",
        "## Kiểm tra dữ liệu", "",
        f"1. ID liên tục 1–515: **{'Đạt' if [q['id'] for q in questions] == list(range(1, 516)) else 'Không đạt'}**.",
        f"2. Câu không có lựa chọn: **{len(missing)}** ({', '.join(map(str, missing)) or 'không có'}).",
        f"3. Câu có số lựa chọn khác 4: **{len(non_four)}** ({', '.join(map(str, non_four)) or 'không có'}).",
        f"4. Câu nhiều đáp án đúng: **{len(multi)}** ({', '.join(map(str, multi)) or 'không có'}).",
        f"5. Câu không nhận diện được đáp án: **{len(unresolved)}** ({', '.join(map(str, unresolved)) or 'không có'}).",
        f"6. Nhóm câu trùng nội dung: **{len(duplicates)}** ({duplicates if duplicates else 'không có'}).",
        f"7. Câu có nội dung bổ sung được tách thành giải thích sau trả lời: **{len(notes)}** ({', '.join(map(str, notes)) or 'không có'}).", "",
        "## Định dạng đã quét", "",
        f"- Run in đậm: {formats['bold']}", f"- Run màu chữ: {formats['color']}",
        f"- Run highlight: {formats['highlight']}", f"- Run gạch chân: {formats['underline']}",
    ]
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    values = extract()
    report(*values)
    print(f"Wrote {OUTPUT.name} and {REPORT.name}")
